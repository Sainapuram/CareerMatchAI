import re

import html

import hashlib

import streamlit as st

from docx import Document

from pypdf import PdfReader

from src.matching import load_occupation_profiles, build_tfidf_model, recommend_careers

# ==========================================

# 1. PAGE CONFIGURATION

# ==========================================

st.set_page_config(

    page_title="CareerMatch AI | Career Advisor",

    page_icon="🎯",

    layout="wide",

    initial_sidebar_state="collapsed"

)

# ==========================================

# 2. CUSTOM STYLING

# ==========================================

st.markdown("""

<style>

    .stApp {

        background-color: #F5F7FB;

    }

    .block-container {

        max-width: 1120px;

        padding-top: 2rem;

        padding-bottom: 3rem;

    }

    h1, h2, h3 {

        color: #142B45;

        font-family: 'Segoe UI', sans-serif;

    }

    .hero {

        background: linear-gradient(120deg, #10243A, #1E4974);

        border-radius: 20px;

        padding: 38px 42px;

        color: white;

        margin-bottom: 28px;

    }

    .hero h1 {

        color: white !important;

        font-size: 2.6rem;

        margin-bottom: 12px;

    }

    .hero p {

        color: #D7E7F6;

        font-size: 1.05rem;

        line-height: 1.7;

    }

    .hero-label {

        color: #A9D6FF;

        font-size: 0.8rem;

        font-weight: 700;

        letter-spacing: 2px;

        margin-bottom: 10px;

    }

    .section-label {

        color: #2563EB;

        font-size: 0.85rem;

        font-weight: 700;

        letter-spacing: 1px;

        margin-top: 8px;

    }

    .section-description {

        color: #64748B;

        font-size: 0.95rem;

        margin-bottom: 14px;

    }

    .stButton > button[kind="primary"] {

        background-color: #2563EB;

        border: none;

        border-radius: 10px;

        color: white;

        font-weight: 600;

        padding: 0.65rem 1.5rem;

    }

    .stButton > button[kind="primary"]:hover {

        background-color: #1D4ED8;

        color: white;

    }

    div[data-testid="stMetric"] {

        background: white;

        border: 1px solid #E2E8F0;

        border-radius: 12px;

        padding: 18px;

    }

    .skill-tag {

        display: inline-block;

        background-color: #EAF2FF;

        color: #1D4ED8;

        padding: 7px 13px;

        border-radius: 20px;

        font-size: 0.88rem;

        font-weight: 600;

        margin: 4px;

    }

    .footer {

        text-align: center;

        color: #64748B;

        font-size: 0.85rem;

        padding-top: 22px;

    }

</style>

""", unsafe_allow_html=True)

# ==========================================

# 3. RESUME PROCESSING FUNCTIONS

# ==========================================

def extract_docx_text(uploaded_file):

    document = Document(uploaded_file)

    paragraphs = [

        paragraph.text

        for paragraph in document.paragraphs

    ]

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                paragraphs.append(cell.text)

    return "\n".join(paragraphs)

def extract_pdf_text(uploaded_file):

    reader = PdfReader(uploaded_file)

    pages = [

        page.extract_text() or ""

        for page in reader.pages

    ]

    return "\n".join(pages)

def remove_email(text):
    pattern = r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
    return re.sub(pattern, "[EMAIL REMOVED]", text)

def remove_phone(text):
    pattern = r"(?<!\w)(?:\+1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}(?!\w)"
    return re.sub(pattern, "[PHONE REMOVED]", text)

def clean_resume_text(text):

    text = remove_email(text)

    text = remove_phone(text)

    return text

# ==========================================

# 4. SKILL EXTRACTION

# ==========================================

SKILLS_LIST = [

    "Python",

    "SQL",

    "Excel",

    "Tableau",

    "Power BI",

    "Machine Learning",

    "Data Analysis",

    "Java",

    "JavaScript",

    "Statistics",

    "Communication",

    "Problem Solving",

    "Project Management",

    "Leadership",

    "Cloud Computing",

    "Deep Learning",

    "Data Visualization",

    "Research",

    "Customer Service",

    "Business Analysis",

    "Data Engineering",

    "Public Speaking",

    "Teamwork",

    "Critical Thinking",

    "Microsoft Office",

    "Git",

    "Docker",

    "AWS",

    "Azure",

    "TensorFlow",

    "PyTorch",

    "FastAPI",

    "Streamlit"

]

def extract_skills(text):

    found_skills = []

    for skill in SKILLS_LIST:

        pattern = r"\b" + re.escape(skill) + r"\b"

        if re.search(pattern, text, re.IGNORECASE):

            found_skills.append(skill)

    return found_skills

# ==========================================

# 5. SESSION STATE

# ==========================================

if "resume_id" not in st.session_state:

    st.session_state.resume_id = None

if "skills_selection" not in st.session_state:

    st.session_state.skills_selection = []

if "custom_skills" not in st.session_state:

    st.session_state.custom_skills = []

if "profile_visible" not in st.session_state:

    st.session_state.profile_visible = False

def add_custom_skill():

    new_skill = st.session_state.custom_skill_input.strip()

    if not new_skill:

        st.session_state.skill_message = (

            "Please enter a skill."

        )

        return

    if len(new_skill) > 60:

        st.session_state.skill_message = (

            "Skill name must be 60 characters or less."

        )

        return

    existing_skills = (

        SKILLS_LIST

        + st.session_state.custom_skills

    )

    matching_skill = next(

        (

            skill for skill in existing_skills

            if skill.casefold() == new_skill.casefold()

        ),

        None

    )

    skill_to_add = matching_skill or new_skill

    if skill_to_add not in st.session_state.custom_skills:

        if skill_to_add not in SKILLS_LIST:

            st.session_state.custom_skills.append(

                skill_to_add

            )

    if skill_to_add not in st.session_state.skills_selection:

        st.session_state.skills_selection.append(

            skill_to_add

        )

    st.session_state.custom_skill_input = ""

    st.session_state.skill_message = (

        "Skill added successfully."

    )

# ==========================================
# 6. HERO HEADER
# ==========================================

st.markdown(
    """
<div class="hero">
    <div class="hero-label">
        ✦ YOUR INTELLIGENT CAREER COMPANION
    </div>
    <h1>CareerMatch AI</h1>
    <p>
        Your next career move starts here.<br>
        Discover your strengths, explore new career paths,
        and build a future that fits your goals.
    </p>
</div>
    """,
    unsafe_allow_html=True
)

# ==========================================

# 7. CAREER GOAL SELECTION

# ==========================================

st.markdown(

    '<div class="section-label">01 / YOUR DIRECTION</div>',

    unsafe_allow_html=True

)

st.subheader("What would you like to achieve?")

st.markdown(

    '<div class="section-description">'

    'Choose the journey that best describes your goals.'

    '</div>',

    unsafe_allow_html=True

)

goal_options = {

    "Explore Career": {

        "icon": "🧭",

        "description": "Discover careers that match your interests."

    },

    "Change Career": {

        "icon": "🔄",

        "description": "Transition into a different profession."

    },

    "Grow Career": {

        "icon": "📈",

        "description": "Advance in your current career."

    }

}

career_goal = st.segmented_control(

    "Choose your career path",

    options=list(goal_options.keys()),

    default="Explore Career",

    selection_mode="single",

    width="stretch"

)

if career_goal is None:

    career_goal = "Explore Career"

goal_details = goal_options[career_goal]

st.info(

    goal_details["icon"] + " " +

    goal_details["description"]

)

# ==========================================

# 8. CAREER PREFERENCES

# ==========================================

st.write("")

st.markdown(

    '<div class="section-label">02 / YOUR AMBITIONS</div>',

    unsafe_allow_html=True

)

st.subheader("Tell us about yourself")

left_col, right_col = st.columns(2, gap="large")

with left_col:

    current_career = ""

    if career_goal in ["Change Career", "Grow Career"]:

        current_career = st.text_input(

            "Current profession",

            placeholder="e.g., Teacher, Software Developer"

        )

    target_career = st.text_input(

        "Target career (optional)",

        placeholder="e.g., Data Analyst"

    )

with right_col:

    experience_years = st.number_input(

        "Years of professional experience",

        min_value=0,

        max_value=50,

        value=0,

        step=1

    )

    learning_hours = st.number_input(

        "Hours available for learning per week",

        min_value=0,

        max_value=80,

        value=5,

        step=1

    )

career_interests = st.text_area(

    "Your interests and career goals",

    placeholder=(

        "Tell us what you enjoy, what you're good at, "

        "and what kind of career you're looking for."

    ),

    height=110

)

# ==========================================

# 9. RESUME UPLOAD

# ==========================================

st.write("")

st.markdown(

    '<div class="section-label">03 / YOUR BACKGROUND</div>',

    unsafe_allow_html=True

)

st.subheader("Upload your résumé")

st.caption(

    "Optional — PDF or DOCX, maximum 5 MB. "

    "You can also enter your skills manually."

)

uploaded_file = st.file_uploader(

    "Upload your résumé",

    type=["pdf", "docx"],

    max_upload_size=5,

    label_visibility="collapsed"

)

extracted_skills = []

resume_text = ""

if uploaded_file is not None:

    if uploaded_file.size > 5 * 1024 * 1024:

        st.error("Maximum file size is 5 MB.")

        st.stop()

    try:

        if uploaded_file.name.lower().endswith(".pdf"):

            resume_text = extract_pdf_text(uploaded_file)

        elif uploaded_file.name.lower().endswith(".docx"):

            resume_text = extract_docx_text(uploaded_file)

        else:

            st.error("Unsupported file format.")

            st.stop()

    except Exception:

        st.error(

            "Unable to process this résumé. "

            "Please upload a valid PDF or DOCX file."

        )

        st.stop()

    if resume_text.strip():

        cleaned_resume = clean_resume_text(resume_text)

        extracted_skills = extract_skills(cleaned_resume)

        st.success(

            "Résumé processed successfully. "

            "Basic email and phone masking applied."

        )

        st.caption(

            f"{len(resume_text):,} characters extracted"

        )

    else:

        st.warning(

            "No readable text was found. "

            "Scanned PDFs may require OCR."

        )

# ==========================================

# 10. SKILL EDITING

# ==========================================

st.write("")

st.markdown(

    '<div class="section-label">04 / YOUR STRENGTHS</div>',

    unsafe_allow_html=True

)

st.subheader("Review and personalize your skills")

st.caption(

    "Confirm extracted skills, remove incorrect matches, "

    "or add skills that are missing."

)

# Identify the uploaded resume

current_resume_id = None

if uploaded_file is not None and resume_text.strip():

    current_resume_id = hashlib.sha256(

        uploaded_file.getvalue()

    ).hexdigest()

# When a different resume is uploaded,

# initialize its extracted skills

if current_resume_id != st.session_state.resume_id:

    st.session_state.resume_id = current_resume_id

    st.session_state.skills_selection = list(

        extracted_skills

    )

# Build the available skills catalog

skill_options = sorted(

    set(SKILLS_LIST)

    | set(st.session_state.custom_skills)

    | set(st.session_state.skills_selection)

)

# One multiselect for both adding and removing skills

st.multiselect(

    "Select all skills you have",

    options=skill_options,

    key="skills_selection",

    help=(

        "Automatically detected skills are preselected. "

        "Add missing skills or remove incorrect ones."

    )

)

st.caption(

    "Skills detected from a résumé are suggestions, "

    "not verified proficiency assessments."

)

# Add custom skills

st.text_input(

    "Add a skill not listed above",

    placeholder="e.g., Healthcare Analytics",

    key="custom_skill_input"

)

st.button(

    "Add Custom Skill",

    on_click=add_custom_skill

)

if "skill_message" in st.session_state:

    st.info(st.session_state.skill_message)

    del st.session_state.skill_message

# Confirmed skills

confirmed_skills = list(

    st.session_state.skills_selection

)

st.markdown("**Your confirmed skills**")

if confirmed_skills:

    skill_html = "".join(

        '<span class="skill-tag">'

        + html.escape(skill)

        + '</span>'

        for skill in confirmed_skills

    )

    st.markdown(

        skill_html,

        unsafe_allow_html=True

    )

    st.caption(

        f"{len(confirmed_skills)} skills selected"

    )

else:

    st.info(

        "No skills selected yet. Upload a résumé "

        "or add skills manually."

    )

# ==========================================

# 11. CAREER PROFILE DASHBOARD

# ==========================================

st.write("")

st.markdown(

    '<div class="section-label">05 / YOUR CAREER PROFILE</div>',

    unsafe_allow_html=True

)

st.subheader("Your personalized career profile")

if st.button(

    "✨ View My Career Profile",

    type="primary",

    use_container_width=True

):

    st.session_state.profile_visible = True

if st.session_state.profile_visible:

    st.write("")

    metric1, metric2, metric3 = st.columns(3)

    with metric1:

        st.metric(

            "Confirmed Skills",

            len(confirmed_skills)

        )

    with metric2:

        st.metric(

            "Years of Experience",

            experience_years

        )

    with metric3:

        st.metric(

            "Learning Hours / Week",

            learning_hours

        )

    st.write("")

    profile_left, profile_right = st.columns(2)

    with profile_left:

        st.markdown("#### Career Direction")

        st.write("**Goal:**", career_goal)

        if current_career.strip():

            st.write(

                "**Current Profession:**",

                current_career

            )

        st.write(

            "**Target Career:**",

            target_career.strip() or "Not specified"

        )

    with profile_right:

        st.markdown("#### Your Strengths")

        if confirmed_skills:

            st.write(", ".join(confirmed_skills))

        else:

            st.write("No skills confirmed yet.")

    if career_interests.strip():

        st.markdown("#### Interests and Goals")

        st.write(career_interests)

    st.info(

        "Career matching is now available. "
        "Skill-gap analysis and personalized learning pathways "
        "are planned for the next development phase."

    )

# ==========================================
# 12. CAREER RECOMMENDATIONS
# ==========================================

st.write("")
st.markdown(
    '<div class="section-label">06 / YOUR CAREER MATCHES</div>',
    unsafe_allow_html=True,
)
st.subheader("Explore your top career matches")
st.caption(
    "Recommendations are ranked by TF-IDF cosine similarity between your "
    "confirmed skills and O*NET occupation profiles. These are not hiring "
    "probabilities or measures of eligibility."
)

@st.cache_resource(show_spinner="Preparing occupation recommendations...")
def get_career_model():
    occupations = load_occupation_profiles()
    return build_tfidf_model(occupations)

if not confirmed_skills:
    st.info("Select at least one skill above to see career recommendations.")
else:
    if st.button("🎯 Find My Top 5 Careers", type="primary", use_container_width=True):
        try:
            profiles, vectorizer, occupation_matrix = get_career_model()
            matches = recommend_careers(
                confirmed_skills,
                profiles,
                vectorizer,
                occupation_matrix,
                top_n=5,
            )
            if matches["Similarity Score"].max() <= 0:
                st.warning(
                    "No meaningful text overlap was found for these skills. "
                    "Try adding more specific skills."
                )
            else:
                display_matches = matches.rename(columns={
                    "O*NET-SOC Code": "O*NET Code",
                    "Title": "Career",
                    "Similarity Score": "Cosine Similarity",
                }).copy()
                display_matches.insert(0, "Rank", range(1, len(display_matches) + 1))
                st.dataframe(
                    display_matches,
                    hide_index=True,
                    use_container_width=True,
                    column_config={
                        "Cosine Similarity": st.column_config.NumberColumn(format="%.4f")
                    },
                )
                if (matches["Similarity Score"] == 0).any():
                    st.caption(
                        "Some occupations have zero text similarity; "
                        "they should not be interpreted as meaningful matches."
                    )
                st.caption(
                    "Career mode, experience, interests, and target career are "
                    "captured in your profile but do not yet affect the ranking."
                )
        except (FileNotFoundError, ImportError, KeyError, ValueError, TypeError) as exc:
            st.error(
                "Career recommendations could not be loaded. Check that "
                "data/processed/occupation_profiles.pkl and src/matching.py "
                "are present and compatible."
            )
            st.caption(f"Technical detail: {exc}")

# ==========================================

# 13. FOOTER

# ==========================================

st.write("")

st.divider()

st.markdown("""

<div class="footer">

    CareerMatch AI · Career Exploration ·

    Career Transitions · Professional Growth

</div>

""", unsafe_allow_html=True)
